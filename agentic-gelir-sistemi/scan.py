"""Periyodik lead taraması — tek çevrim.

Her çağrıldığında: ilanları skorlar (lead_research), sonuçları zaman damgalı olarak
reports/lead_scan.xlsx dosyasına EKLER (üzerine yazmaz — geçmiş birikir) ve özet basar.
Arka plan döngüsü bu betiği belirli aralıklarla çağırır, çıktıyı gösterir ve Excel'i git'e pushlar.

Çalıştır:  python3 scan.py [data/live_listings.json]
Maliyet:   0₺  (LLM_MODE=mock — henüz gerçek API yok)
"""
import os
import sys
import json
from datetime import datetime, timezone, timedelta

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from agents import lead_research  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402

ROOT = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(ROOT, "reports", "lead_scan.xlsx")
DEFAULT_DATA = os.path.join(ROOT, "data", "live_listings.json")
HEADERS = ["taranma_zamani", "kaynak", "ilan", "skor", "neden", "eslesme", "link", "durum"]
LINK_COL = 7  # 'link' sütununun 1-tabanlı indeksi (Excel'de köprü için)


TR = timezone(timedelta(hours=3))  # Türkiye saati (UTC+3, kalıcı — yaz saati yok)


def _now():
    return datetime.now(TR).strftime("%Y-%m-%d %H:%M:%S TRT")


def _ensure_wb():
    os.makedirs(os.path.dirname(XLSX), exist_ok=True)
    if os.path.exists(XLSX):
        wb = load_workbook(XLSX)
        return wb, wb.active
    wb = Workbook()
    ws = wb.active
    ws.title = "leads"
    ws.append(HEADERS)
    return wb, ws


def main():
    data_path = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_DATA
    with open(data_path, encoding="utf-8") as f:
        listings = json.load(f)

    scored = sorted(lead_research.run(listings), key=lambda x: -x["score"])
    ts = _now()

    wb, ws = _ensure_wb()
    for s in scored:
        url = s.get("url", "")
        ws.append([ts, s.get("source", "?"), s["title"], s["score"],
                   s["reason"], ", ".join(s["skills"]), url, "queued"])
        if url:  # link hücresini tıklanabilir köprü yap
            cell = ws.cell(row=ws.max_row, column=LINK_COL)
            cell.hyperlink = url
            cell.style = "Hyperlink"
    wb.save(XLSX)

    top = [s for s in scored if s["score"] >= 90]
    print(f"\n🔁 [{ts}] Tarama tamamlandı — {len(scored)} ilan işlendi.")
    print(f"   Excel güncellendi: reports/lead_scan.xlsx (+{len(scored)} satır)")
    print(f"   Bu taramada {len(top)} tam-niş (skor 90) lead:\n")
    for s in scored:
        flag = "⭐" if s["score"] >= 90 else "  "
        print(f"   {flag} [{s['score']:>3}] {s['title']}  ({s.get('source')})")
        print(f"          {s.get('url', '(link yok)')}")
    print()


if __name__ == "__main__":
    main()
